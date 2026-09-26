"""Inspect layout of committed government spreadsheets without guessing schemas.

Prints worksheet names, dimensions and first nonempty rows. This is also
useful locally when a government portal changes its report layout.
"""
from __future__ import annotations
import argparse
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DEFAULT = ROOT / "datasets/dilrmp/Uttar Pradesh"


def inspect_workbook(path: Path, sample: int = 14):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        print(f"FILE: {path.relative_to(ROOT) if path.is_relative_to(ROOT) else path}")
        for sheet in book:
            print(f"SHEET: {sheet.title!r} dimensions={sheet.max_row}x{sheet.max_column}")
            nonempty = 0
            for line_no, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = [(i + 1, str(v).strip()[:100]) for i, v in enumerate(row) if v is not None and str(v).strip()]
                if not cells:
                    continue
                print(f"  ROW {line_no}: {cells[:24]}")
                nonempty += 1
                if any('GHAZIABAD' in str(v).upper() for v in row if v is not None):
                    print(f'  TARGET GHAZIABAD ROW {line_no}: {cells[:40]}')
                if nonempty >= sample and line_no > 100:
                    break
                if nonempty >= sample:
                    # Still scan for Ghaziabad, without printing other district rows.
                    for other_line, other in enumerate(sheet.iter_rows(min_row=line_no+1, values_only=True), start=line_no+1):
                        if any('GHAZIABAD' in str(v).upper() for v in other if v is not None):
                            match = [(i+1,str(v).strip()[:100]) for i,v in enumerate(other) if v is not None]
                            print(f'  TARGET GHAZIABAD ROW {other_line}: {match[:40]}')
                    break
    finally:
        book.close()


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--sample", type=int, default=14)
    p.add_argument("--dir", type=Path, default=DEFAULT)
    args = p.parse_args()
    for file in sorted(args.dir.rglob("*.xlsx")):
        inspect_workbook(file, args.sample)
