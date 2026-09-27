"""Regression tests for checked document INSERT and XLSX parsing.

Load only selected AST functions from main.py so tests don't download
sentence-transformer models or require running PostgreSQL.
"""
from __future__ import annotations

import ast
import io
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

MAIN = Path(__file__).resolve().parents[1] / "app" / "main.py"


def isolated(*names):
    tree = ast.parse(MAIN.read_text(encoding="utf-8"))
    selected = [node for node in tree.body if isinstance(node, ast.FunctionDef)
                and node.name in names]
    assert {node.name for node in selected} == set(names)
    module = ast.Module(
        body=[ast.ImportFrom(module="__future__",
                             names=[ast.alias(name="annotations")], level=0), *selected],
        type_ignores=[],
    )
    namespace = {"Path": Path, "io": io, "load_workbook": load_workbook,
                 "HTTPException": HTTPException}
    exec(compile(ast.fix_missing_locations(module), str(MAIN), "exec"),
         namespace)
    return namespace


class FakeCursor:
    def __init__(self, row):
        self.row = row

    def fetchone(self):
        return self.row


class FakeDb:
    def __init__(self, row):
        self.row = row
        self.statements = []

    def execute(self, query, params):
        self.statements.append((query, params))
        return FakeCursor(self.row)


def test_insert_returns_integer_and_preserves_metadata_in_same_transaction():
    fn = isolated("insert_document")["insert_document"]
    db = FakeDb((42,))
    result = fn(db, "source.pdf", "digest", 8,
                ("Report", "Uttar Pradesh", "", "https://example.org"))
    assert result == 42
    assert len(db.statements) == 2
    assert db.statements[0][1] == ("source.pdf", "digest", 8)
    assert db.statements[1][1][-1] == 42


def test_insert_with_none_result_raises_clear_error():
    fn = isolated("insert_document")["insert_document"]
    with pytest.raises(RuntimeError, match="did not return an ID"):
        fn(FakeDb(None), "document.pdf", "digest", 1)


def test_upload_xlsx_extracts_multiple_worksheets_as_separate_pages():
    fn = isolated("extract_pages")["extract_pages"]
    book = Workbook()
    first = book.active
    first.title = "First"
    first.append(["District", "Count"])
    first.append(["Ghaziabad", 264])
    book.create_sheet("Second").append(["Other", "Evidence"])
    buffer = io.BytesIO()
    book.save(buffer)
    pages = fn("report.xlsx", buffer.getvalue())
    assert len(pages) == 2
    assert "Ghaziabad | 264" in pages[0][1]
    assert "worksheet: Second" in pages[1][1]


def test_xlsx_limit_rejects_excessive_nonempty_rows():
    fn = isolated("extract_pages")["extract_pages"]
    book = Workbook()
    sheet = book.active
    for _ in range(10_001):
        sheet.append(["a"])
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(HTTPException) as caught:
        fn("excess.xlsx", buffer.getvalue())
    assert caught.value.status_code == 413
